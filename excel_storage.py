from io import BytesIO
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook


VEHICLE_SHEET_NAMES = ("vehiculo", "vehiculos")


def _normalize(value) -> str:
    normalized = unicodedata.normalize("NFD", str(value).strip().lower())
    return "".join(char for char in normalized if unicodedata.category(char) != "Mn")


def _find_sheet(sheet_names, candidates):
    by_normalized_name = {name.strip().lower(): name for name in sheet_names}
    for candidate in candidates:
        if candidate in by_normalized_name:
            return by_normalized_name[candidate]
    return None


def read_reference_sheets(workbook_bytes: bytes):
    """Read vehicle and driver tables from supported workbook sheets."""
    workbook = pd.ExcelFile(BytesIO(workbook_bytes))
    vehicle_sheet = _find_sheet(workbook.sheet_names, VEHICLE_SHEET_NAMES)
    driver_sheet = _find_sheet(workbook.sheet_names, ("conductores",))
    if vehicle_sheet is None:
        raise ValueError("No se encuentra la pestaña 'vehiculo' o 'vehiculos'.")
    if driver_sheet is None:
        raise ValueError("No se encuentra la pestaña 'conductores'.")

    vehicles = workbook.parse(vehicle_sheet)
    drivers = workbook.parse(driver_sheet)
    workbook.close()
    if vehicles.shape[1] < 5:
        raise ValueError(
            "La pestaña de vehículos debe tener el tipo de vehículo en la columna B "
            "y la matrícula del vehículo en la columna E."
        )
    return vehicles, drivers


def _split_formula_parts(expression: str, separators: str) -> list[str]:
    parts = []
    start = 0
    inside_quotes = False
    nesting = 0
    index = 0
    while index < len(expression):
        char = expression[index]
        if char == '"':
            if inside_quotes and index + 1 < len(expression) and expression[index + 1] == '"':
                index += 2
                continue
            inside_quotes = not inside_quotes
        elif not inside_quotes:
            if char == "(":
                nesting += 1
            elif char == ")":
                nesting -= 1
            elif char in separators and nesting == 0:
                parts.append(expression[start:index].strip())
                start = index + 1
        index += 1
    parts.append(expression[start:].strip())
    return parts


def _formula_cell_value(address, values_sheet, formulas_sheet, depth):
    if depth > 12:
        raise ValueError("La fórmula de matrícula tiene referencias circulares.")

    address = address.replace("$", "")
    value = values_sheet[address].value
    if value is not None:
        return str(value)

    formula = formulas_sheet[address].value
    if isinstance(formula, str) and formula.startswith("="):
        return _evaluate_concat_formula(
            formula, values_sheet, formulas_sheet, depth + 1
        )
    return ""


def _formula_token_value(token, values_sheet, formulas_sheet, depth):
    token = token.strip()
    if len(token) >= 2 and token[0] == '"' and token[-1] == '"':
        return token[1:-1].replace('""', '"')

    if "!" in token:
        token = token.rsplit("!", 1)[1]

    if re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+", token, flags=re.IGNORECASE):
        return _formula_cell_value(token, values_sheet, formulas_sheet, depth + 1)

    range_match = re.fullmatch(
        r"(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)",
        token,
        flags=re.IGNORECASE,
    )
    if range_match:
        start = range_match.group(1).replace("$", "")
        end = range_match.group(2).replace("$", "")
        from openpyxl.utils.cell import range_boundaries

        min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}")
        return "".join(
            _formula_cell_value(
                values_sheet.cell(row=row, column=column).coordinate,
                values_sheet,
                formulas_sheet,
                depth + 1,
            )
            for row in range(min_row, max_row + 1)
            for column in range(min_col, max_col + 1)
        )

    if "&" in token or re.match(
        r"(?:_xlfn\.)?(?:CONCATENAR|CONCATENATE|CONCAT)\(",
        token,
        flags=re.IGNORECASE,
    ):
        return _evaluate_concat_formula(
            f"={token}", values_sheet, formulas_sheet, depth + 1
        )

    return token


def _evaluate_concat_formula(formula, values_sheet, formulas_sheet, depth=0):
    if depth > 12:
        raise ValueError("La fórmula de matrícula tiene referencias circulares.")

    expression = formula.lstrip("=").strip()
    function_match = re.fullmatch(
        r"(?:_xlfn\.)?(?:CONCATENAR|CONCATENATE|CONCAT)\((.*)\)",
        expression,
        flags=re.IGNORECASE,
    )
    if function_match:
        parts = _split_formula_parts(function_match.group(1), ",;")
    elif "&" in expression:
        parts = _split_formula_parts(expression, "&")
    else:
        raise ValueError("La fórmula de matrícula no es una concatenación compatible.")

    return "".join(
        _formula_token_value(part, values_sheet, formulas_sheet, depth + 1)
        for part in parts
    )


def read_vehicle_registrations(workbook_bytes: bytes, selected_type: str) -> list[str]:
    """Read column E registrations whose column B vehicle type matches selection."""
    formulas_book = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=True)
    values_book = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    vehicle_sheet = _find_sheet(formulas_book.sheetnames, VEHICLE_SHEET_NAMES)
    if vehicle_sheet is None:
        raise ValueError("No se encuentra la pestaña 'vehiculo' o 'vehiculos'.")

    formulas_sheet = formulas_book[vehicle_sheet]
    values_sheet = values_book[vehicle_sheet]
    registrations = []
    for row in range(2, 26):
        vehicle_type = values_sheet.cell(row=row, column=2).value
        if _normalize(vehicle_type) != selected_type:
            continue

        registration = values_sheet.cell(row=row, column=5).value
        formula = formulas_sheet.cell(row=row, column=5).value
        if (registration is None or registration == "") and isinstance(formula, str):
            if formula.startswith("="):
                registration = _evaluate_concat_formula(
                    formula, values_sheet, formulas_sheet
                )

        if registration is not None and str(registration).strip():
            registrations.append(str(registration).strip())

    formulas_book.close()
    values_book.close()
    return list(dict.fromkeys(registrations))


def _normalized_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", _normalize(value))


def _header_index(headers, field_name):
    normalized = [_normalized_header(header) for header in headers]
    field = _normalized_header(field_name)
    aliases = {
        "vehiculo": ("matriculamodelo", "vehiculo"),
    }
    for candidate in aliases.get(field, (field,)):
        if candidate in normalized:
            return normalized.index(candidate)
    return None


def append_record(workbook_bytes: bytes, record: dict) -> bytes:
    """Write one record after the last populated row in 'base datos'."""
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet_name = _find_sheet(workbook.sheetnames, ("base datos",))
    if sheet_name is None:
        raise ValueError("El Excel no contiene la hoja 'base datos'.")

    sheet = workbook[sheet_name]
    headers = [
        sheet.cell(row=1, column=col).value
        for col in range(1, sheet.max_column + 1)
    ]
    while headers and headers[-1] is None:
        headers.pop()

    if not any(headers):
        headers = list(record)
        for column, header in enumerate(headers, start=1):
            sheet.cell(row=1, column=column, value=header)

    for field_name in record:
        if _header_index(headers, field_name) is None:
            headers.append(field_name)
            sheet.cell(row=1, column=len(headers), value=field_name)

    last_data_row = 1
    for row in range(sheet.max_row, 1, -1):
        if any(
            sheet.cell(row=row, column=col).value is not None
            for col in range(1, max(sheet.max_column, len(headers)) + 1)
        ):
            last_data_row = row
            break

    # Move earlier vehicle values out of a duplicate "vehiculo" column (e.g. M)
    # into the existing MATRICULA-MODELO column (D), without overwriting data.
    normalized_headers = [_normalized_header(header) for header in headers]
    vehicle_destination = next(
        (index + 1 for index, name in enumerate(normalized_headers)
         if name == "matriculamodelo"),
        None,
    )
    vehicle_source = next(
        (index + 1 for index, name in enumerate(normalized_headers)
         if name == "vehiculo"),
        None,
    )
    if vehicle_destination and vehicle_source and vehicle_destination != vehicle_source:
        for row in range(2, last_data_row + 1):
            source_cell = sheet.cell(row=row, column=vehicle_source)
            destination_cell = sheet.cell(row=row, column=vehicle_destination)
            if source_cell.value not in (None, "") and destination_cell.value in (None, ""):
                destination_cell.value = source_cell.value
                source_cell.value = None

    new_row = last_data_row + 1
    row_values = [None] * len(headers)
    for field_name, value in record.items():
        column_index = _header_index(headers, field_name)
        if column_index is not None:
            row_values[column_index] = value
    for column, value in enumerate(row_values, start=1):
        sheet.cell(row=new_row, column=column, value=value)

    for table in sheet.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        max_row = max(max_row, new_row)
        table.ref = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        )
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref

    if sheet.auto_filter.ref:
        min_col, min_row, max_col, max_row = range_boundaries(sheet.auto_filter.ref)
        sheet.auto_filter.ref = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max(max_col, len(headers)))}{max(max_row, new_row)}"
        )

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
